import openpyxl

from fnsportal.sources import tablestream as ts


def test_html_stream_with_cp1251_colspan_and_entities(tmp_path):
    path = tmp_path / "t.html"
    path.write_bytes("""<html><head><meta charset="windows-1251"></head><table>
      <tr><th>Код региона</th><th colspan="2">Наименование</th></tr>
      <tr><td>77</td><td>&laquo;Москва&raquo;</td><td><b>2,2&nbsp;%</b></td></tr>
      <tr><td>50</td><td>Московская<br/>область</td><td>1,5 %</td></tr>
    </table></html>""".encode("cp1251"))
    rows = list(ts.iter_rows(path))
    assert rows[0] == ["Код региона", "Наименование", ""]
    assert rows[1] == ["77", "«Москва»", "2,2 %"]
    assert rows[2][1] == "Московская область"


def test_csv_detects_encoding_and_delimiter(tmp_path):
    path = tmp_path / "t.csv"
    path.write_bytes("код;наименование\n77;Москва\n".encode("cp1251"))
    assert list(ts.iter_rows(path)) == [["код", "наименование"], ["77", "Москва"]]


def test_header_detection_prefers_real_header(tmp_path):
    path = tmp_path / "t.csv"
    path.write_text(
        "Отчёт;;\n"
        "Код региона;Вид налога;Размер ставки;Вид льготы\n"
        "77;Транспортный налог;12;Пониженная ставка\n", encoding="utf-8")
    idx, header, _ = ts.find_header(ts.iter_rows(path),
                                    required=["налог", "регион", "ставка", "льгот"])
    assert idx == 1 and header[0] == "Код региона"


def test_xlsx_sheets_are_streamed_separately(tmp_path):
    path = tmp_path / "t.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Лист1"
    wb.active.append(["Показатель", "Код"])
    wb.create_sheet("Лист2").append(["a", "b"])
    wb.save(path)
    names = [(table.name, list(table.rows)) for table in ts.iter_tables(path)]
    assert [name for name, _ in names] == ["Лист1", "Лист2"]
    assert names[0][1][0] == ["Показатель", "Код"]
